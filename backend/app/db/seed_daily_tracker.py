import openpyxl
from datetime import datetime
from app.db.database import engine, Base, SessionLocal
from app.models.user import User
from app.models.daily_tracker import DailyTracker

def normalize_date(v):
    if not v:
        return None
    s = str(v).strip()
    if s.lower() in ["weekend", "none", ""]:
        return None
    if isinstance(v, datetime):
        # Month and day were transposed by Excel when parsed with US locale
        return f"{v.year}-{str(v.day).zfill(2)}-{str(v.month).zfill(2)}"
    if "/" in s:
        parts = s.split("/")
        if len(parts) == 3:
            d, m, y = parts
            return f"{y}-{m.zfill(2)}-{d.zfill(2)}"
    return s[:10]

def seed():
    # Ensure tables are created
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    # Load users for mapping
    users = db.query(User).all()
    user_map = {}
    for u in users:
        user_map[u.name.lower()] = u
        # Add first name mapping
        first = u.name.lower().split()[0]
        user_map[first] = u

    # Load workbook
    excel_path = r"c:\Users\saathithy\Downloads\C2L_Daily Update Tracker.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["Sheet1"]

    engineers = []
    for col in range(2, 29, 3):
        name = sheet.cell(row=1, column=col).value
        if name:
            engineers.append({"name": str(name).strip(), "start_col": col})

    # Clear existing entries in c2l_daily_tracker to avoid duplicates on re-run
    db.query(DailyTracker).delete()
    db.commit()

    created_count = 0
    for r in range(3, sheet.max_row + 1):
        raw_date = sheet.cell(row=r, column=1).value
        parsed_d = normalize_date(raw_date)
        if not parsed_d:
            continue

        for eng in engineers:
            p1 = sheet.cell(row=r, column=eng["start_col"]).value
            p2 = sheet.cell(row=r, column=eng["start_col"] + 1).value
            st = sheet.cell(row=r, column=eng["start_col"] + 2).value

            p1_str = str(p1).strip() if p1 is not None else ""
            p2_str = str(p2).strip() if p2 is not None else ""
            st_str = str(st).strip() if st is not None else ""

            # Check if there is any content
            if not p1_str and not p2_str and not st_str:
                continue

            # Determine status
            status_val = st_str
            if not status_val:
                if "leave" in p1_str.lower() or "leave" in p2_str.lower():
                    status_val = "Leave"
                elif p1_str or p2_str:
                    status_val = "Completed"
                else:
                    status_val = "In-Progress"

            # Match user
            eng_name_lower = eng["name"].lower()
            matched_user = user_map.get(eng_name_lower)
            if not matched_user:
                for k, u in user_map.items():
                    if k in eng_name_lower or eng_name_lower in k:
                        matched_user = u
                        break

            user_id = matched_user.id if matched_user else None
            official_name = matched_user.name if matched_user else eng["name"]

            tracker = DailyTracker(
                user_id=user_id,
                user_name=official_name,
                task_date=parsed_d,
                phase_1=p1_str or None,
                phase_2=p2_str or None,
                status=status_val,
            )
            db.add(tracker)
            created_count += 1

    db.commit()
    print(f"Successfully seeded {created_count} daily tracker records into c2l_daily_tracker!")
    db.close()

if __name__ == "__main__":
    seed()
