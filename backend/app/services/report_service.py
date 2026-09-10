import io
import csv
from typing import Optional
from sqlalchemy.orm import Session
from sqlalchemy import desc
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.c2l_batch import C2LBatch
from app.models.user import User

def generate_report_excel(
    db: Session,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    status: Optional[str] = None,
    batch_type: Optional[str] = None,
    assigned_to_id: Optional[int] = None
) -> io.BytesIO:
    query = db.query(C2LBatch)
    if start_date:
        query = query.filter(C2LBatch.start_date >= start_date)
    if end_date:
        query = query.filter(C2LBatch.end_date <= end_date)
    if status and status != "ALL":
        query = query.filter(C2LBatch.work_status == status)
    if batch_type and batch_type != "ALL":
        query = query.filter(C2LBatch.batch_type == batch_type)
    if assigned_to_id:
        query = query.filter(C2LBatch.assigned_to_id == assigned_to_id)
        
    batches = query.order_by(desc(C2LBatch.updated_at)).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "C2L Operations Report"

    # Deep Navy Theme Headers
    navy_fill = PatternFill(start_color="0F2142", end_color="0F2142", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "Sl No.", "Batch No.", "Batch Type", "Location", "Complexity", 
        "Assigned Owner", "Start Date", "End Date", "Total Hours", 
        "Work Status", "QC Status", "Audit Status", "Client Ready?", "Remarks"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for idx, b in enumerate(batches, start=1):
        owner_name = b.assigned_to.name if b.assigned_to else "Unassigned"
        client_ready = "YES" if (b.work_status == "COMPLETED" and b.audit_status == "PASSED") else "NO"
        
        row_data = [
            idx,
            b.batch_no,
            b.batch_type or "",
            b.location or "",
            b.complexity or "",
            owner_name,
            b.start_date or "",
            b.end_date or "",
            b.total_hours or "",
            b.work_status,
            b.qc_status,
            b.audit_status,
            client_ready,
            b.current_remarks or ""
        ]
        ws.append(row_data)
        row_idx = idx + 1
        ws.row_dimensions[row_idx].height = 20
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 7, 8, 9, 10, 11, 12, 13):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column auto-width adjustment
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_report_csv(
    db: Session,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    status: Optional[str] = None,
    batch_type: Optional[str] = None,
    assigned_to_id: Optional[int] = None
) -> io.StringIO:
    query = db.query(C2LBatch)
    if start_date:
        query = query.filter(C2LBatch.start_date >= start_date)
    if end_date:
        query = query.filter(C2LBatch.end_date <= end_date)
    if status and status != "ALL":
        query = query.filter(C2LBatch.work_status == status)
    if batch_type and batch_type != "ALL":
        query = query.filter(C2LBatch.batch_type == batch_type)
    if assigned_to_id:
        query = query.filter(C2LBatch.assigned_to_id == assigned_to_id)
        
    batches = query.order_by(desc(C2LBatch.updated_at)).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Sl No.", "Batch No", "Batch Type", "Location", "Complexity", 
        "Assigned Owner", "Start Date", "End Date", "Total Hours", 
        "Work Status", "QC Status", "Audit Status", "Client Ready", "Remarks"
    ])

    for idx, b in enumerate(batches, start=1):
        owner_name = b.assigned_to.name if b.assigned_to else "Unassigned"
        client_ready = "YES" if (b.work_status == "COMPLETED" and b.audit_status == "PASSED") else "NO"
        writer.writerow([
            idx,
            b.batch_no,
            b.batch_type or "",
            b.location or "",
            b.complexity or "",
            owner_name,
            b.start_date or "",
            b.end_date or "",
            b.total_hours or "",
            b.work_status,
            b.qc_status,
            b.audit_status,
            client_ready,
            b.current_remarks or ""
        ])

    output.seek(0)
    return output
